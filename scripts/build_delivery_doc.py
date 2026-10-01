"""Generate the editable, code-and-evidence-bounded delivery handbook."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.68), Inches(.65)
s.left_margin = s.right_margin = Inches(.75)
s.header_distance = s.footer_distance = Inches(.28)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.14
doc.styles['Title'].font.size = Pt(24)
doc.styles['Subtitle'].font.size = Pt(13)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
for n in ('Heading 1', 'Heading 2'):
    doc.styles[n].paragraph_format.space_before = Pt(9)
    doc.styles[n].paragraph_format.keep_with_next = True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
header = s.header.paragraphs[0]
header.text = '唐岛湾北岸声学边缘交通演示  项目交付手册'
header.runs[0].font.size = Pt(8)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('交付 1.2.0  |  2026年10月1日  |  ')
for field in ('PAGE', 'NUMPAGES'):
    if field == 'NUMPAGES': footer.add_run(' / ')
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), field); footer._p.append(f)
for run in footer.runs: run.font.size = Pt(8)
doc.core_properties.title = '唐岛湾声学边缘交通演示项目交付手册'
doc.core_properties.subject = '运行操作 系统设计 能力边界 验证与移机'
doc.core_properties.author = ''

def p(text, bold=False, size=None):
    x = doc.add_paragraph(); run = x.add_run(text); run.bold = bold
    if size: run.font.size = Pt(size)
    return x

def h(text):
    return doc.add_heading(text, level=2)

def page(title):
    x = doc.add_heading(title, level=1)
    x.paragraph_format.page_break_before = True

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    for c, width in zip(t.columns, widths): c.width = Inches(width)
    for i, values in enumerate([headers, *rows]):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        tr = cells[0]._tc.getparent(); trpr = tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for cell, text, width in zip(cells, values, widths):
            cell.width = Inches(width); cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EFF5' if i == 0 else 'FFFFFF'); pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                item = OxmlElement('w:' + edge); item.set(qn('w:w'), '100'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            pr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'bottom', 'left', 'right'):
                item = OxmlElement('w:' + edge); item.set(qn('w:val'), 'single'); item.set(qn('w:sz'), '4'); item.set(qn('w:color'), 'D1D8DE'); borders.append(item)
            pr.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    run.font.size = Pt(10); run.bold = i == 0
    gap = doc.add_paragraph(); gap.paragraph_format.space_after = Pt(2); gap.paragraph_format.line_spacing = Pt(2)
    return t

doc.add_paragraph('唐岛湾声学边缘交通演示', 'Title')
doc.add_paragraph('项目交付手册  版本 1.2', 'Subtitle')
p('交付标记 1.2.0 公开便携版    2026年10月1日', size=10)
p('本项目已实现真实地图上的三维仿真闭环：车辆经过路口，触发合成声学任务，任务在路侧节点排队计算或跨站卸载，结果回传后驱动有界信号响应。用户可以暂停观察、切换简单调度策略、调整视角并导出记录。')
p('本手册供演示、维护和技术评审使用，说明如何运行、系统如何工作、哪些能力已经验证，以及哪些部分仍需要真实数据和设备。本版提供Windows便携启动，沿用已验证的性能优化代码；道路、车流、任务算法和统计口径保持不变。')
doc.add_picture(str(ROOT / 'evidence/performance/live_scene.png'), width=Inches(7))
p('图1  性能优化版实际页面截图  截图时仿真暂停于601.4秒', size=9)
h('阅读顺序')
p('先操作：第二章；看地图与数据边界：第三章；理解系统和算法：第四、五章；核对验证与性能：第六、七章；评估后续难度：第八章；交付与移机：第九章；复现与发布核对：第十章。')
p('核心边界：道路和建筑轮廓来自真实地图；车辆需求、声学事件、RSU部署、计算与通信耗时是模拟配置。本项目尚未执行真实录音识别，也未证明真实交通效率提升。', bold=True)

page('第二章 启动和地图操作')
h('第一次运行')
p('建议从GitHub Releases下载Windows便携包，完整解压到可写目录，再双击Start_Demo.cmd。便携包包含运行依赖；源码包会在缺少依赖时联网下载固定版本到项目内。等待本地页面打开后点击“继续”，先用1×和轻量画质。首次启动默认暂停。')
p('默认端口为8765，被其他应用或另一份项目占用时依次尝试8766至8775。以本次启动器打开的地址为准；本目录的 runs/server.json 保存实际端口。不要直接双击 web/index.html，网页需要本地HTTP接口。')
h('鼠标和视角快捷操作')
table(['目的', '操作'], [
    ['平移道路位置', '在三维区域按住鼠标右键拖动，上下左右移动地图'],
    ['旋转方向和俯仰', '在三维区域按住鼠标左键拖动'],
    ['放大或缩小', '滚轮向上拉近，向下拉远'],
    ['恢复路口视角', '点击“聚焦路口”；仅移动镜头，不重置仿真'],
    ['查看完整区域或节点', '点击“全景”，或点击右侧某个RSU卡片'],
], [1.65, 5.35])
h('控制与记录')
table(['控件', '实际影响'], [
    ['暂停或继续', '暂停或恢复后端仿真时钟；车辆、任务、信号一起暂停'],
    ['0.5×至4×', '改变仿真推进速度，不改变仿真步长和服务参数'],
    ['重置或切换策略', '开始新的同种子运行，当前指标清零'],
    ['导出实验记录', '保存当前运行的配置、完整任务、事件和指标JSON'],
], [1.65, 5.35])
p('结束时使用 Stop_Demo.cmd 正常停止本目录服务及其SUMO。关闭或隐藏浏览器只影响前端显示，不会暂停正在运行的后端。更新过项目文件后，在旧页面按Ctrl+F5重新加载。')

page('第三章 地图范围和真实性')
p('首版区域按用户确认选择唐岛湾北岸，重点展示漓江西路沿线太行山路、井冈山路、武夷山路和阿里山路附近四个路口。原始OSM快照于2026年10月1日下载并保留，可按来源清单与哈希核对。')
table(['数据', '数量或处理', '解释'], [
    ['道路网络', '324条非内部边', '渲染含转向连接段，共1041条记录，不是1041条独立街道'],
    ['建筑轮廓', '129个', '10个高度取自OSM；7个按楼层×3.2m换算'],
    ['缺失建筑高度', '112个按15m示意', '不是测绘结果，不代表实际楼高'],
    ['水域和公园', '26个区域', '由OSM轮廓生成，无写实地形和水文模型'],
    ['路侧节点', '4个RSU，覆盖半径105m', '部署、外观、比例与可感知范围均是演示配置'],
], [1.1, 2.0, 3.9])
h('真实来源与模拟过程')
p('真实来源指地理轮廓。车道细节、限速或路宽可能包含导入推断；信号配时与车流需求为合成配置，不是当地交管部门实时数据。道路高程与建筑立面尚未建立精确模型，不能将当前场景当作完整数字孪生。')
p('车辆位置、速度、车道和灯态来自正在运行的SUMO。麦克风和机柜为程序生成的简化几何；圆形覆盖和波纹表达任务触发与显示效果，没有求解噪声传播、反射、遮挡、混响或多声源分离。')
h('坐标和来源核对')
p('后端使用SUMO本地平面米制坐标。前端把平面x映射为三维x、平面y映射为负z，高度使用三维y；统一减去场景中心改善数值稳定性。经纬度边界、投影和转换过程见MAP_NOTES.md。')
p('实际数据以data/scene.json为准；下载元数据见data/source_manifest.json，原始地图见data/tangdao.osm.xml，重建命令见data/build_commands.json。CONTRACT.md中的坐标、覆盖半径等JSON值只是结构示例。', size=10)
p('地图 © OpenStreetMap contributors，采用ODbL 1.0。Three.js 0.180.0使用MIT许可；相关署名与许可证随包保留。')

page('第四章 系统结构与接口')
p('离线准备阶段把OSM快照转换为SUMO路网、车流配置和三维场景。运行阶段由Python控制SUMO推进，同时维护感知事件、计算队列与信号响应；浏览器读取快照并呈现场景。')
table(['模块', '入口', '职责'], [
    ['地图构建', 'scripts/build_map.py', '生成网络、需求、场景与来源记录'],
    ['交通引擎', 'scenario/tangdao.sumocfg', 'SUMO推进车辆、车道连接和信号相位'],
    ['调度与控制', 'simulation.py', '任务创建、资源预约、回传、控制检查'],
    ['本地服务', 'server.py', 'HTTP状态、控制、记录导出与正常关闭'],
    ['三维显示', 'web/app.js', '道路与设备几何、车辆插值、事件动画'],
], [1.1, 2.3, 3.6])
h('一次任务的完整顺序')
p('受控进口车辆进入覆盖区 → 合成感知任务 → 选择RSU → 本地发送或跨站传输 → 排队 → 计算 → 回传 → 在下一仿真步兑现完成 → 检查信号响应。事件记录保留实际触发点，中间动画按任务时间戳和仿真时钟推导。')
table(['接口', '用途'], [
    ['GET /api/scene', '道路、建筑、RSU及来源数据'],
    ['GET /api/state', '当前车辆、信号、任务子集、事件和完整累计指标'],
    ['POST /api/control', 'pause、resume、reset、speed、scheduler'],
    ['GET /api/health', '本地服务、SUMO连接与状态'],
    ['GET /api/export', '完整运行配置、任务、事件和统计'],
    ['POST /api/shutdown', '保存并关闭本地服务及其SUMO实例'],
], [2.4, 4.6])
p('HTTP仅绑定本机地址。前端不直接控制交通灯；状态锁协调仿真推进和控制指令。状态接口为显示选取任务子集，不能用画面上的特效数量代替完整任务数；统计与导出保留完整任务记录。')

page('第五章 调度参数与信号规则')
h('感知触发与最早计算完成调度')
p('车辆同时位于RSU覆盖范围和其信号灯受控进口车道时创建任务；每轮同一车辆对同一RSU最多触发一次。车道来自SUMO真值，因此本版不能用于报告真实声学车道识别精度。')
p('对每个候选RSU：任务到达时间等于当前仿真时间加发送耗时；计算开始时间取到达时间与节点已预约完成时间的较大者；计算结束时间再加服务耗时。选择结束最早的节点，平手优先本地。该目标比较计算结束时间，不包括结果返回时间。')
p('节点为单服务器、非抢占队列，已派发的资源预约不重排。另提供“本地执行 · 不卸载”对照策略。切换策略会从相同种子重新开始，不混合两种策略的指标。')
table(['参数', '当前值或公式'], [
    ['仿真步长与种子', '0.2秒；随机种子42'],
    ['基础服务耗时', 'RSU_1至RSU_4分别为2.6、1.6、2.3、1.35秒'],
    ['任务工作量系数', '按固定种子取0.9至1.1，乘以基础服务耗时'],
    ['本地发送及返回', '各0.04秒'],
    ['跨站发送及返回', '发送0.28+d/1600秒；返回0.16+d/2400秒'],
], [1.8, 5.2])
p('其中d为两个节点之间的平面距离，单位米。上述时间是合成配置，分母用于构造演示时延，不代表真实无线链路速率或物理声速。实际运行参数保存在导出的configuration中。', size=10)
h('信号响应边界')
p('任务计算结束并回传后才检查控制：对应进口当前为绿灯、没有黄灯、该相位尚未延长，且未超时长上限，才有限延长当前绿灯。每相位至多一次、每次至多4秒，总绿灯不超过55秒；不直接跳转相位。不满足条件的任务仍完成，但记录“保持配时”。')
p('目前不重新确认车辆当前位置、具体转向或结果新鲜度。有界规则用于仿真展示，不构成真实道路安全认证；真实信号机接入还需要独立验证。')

page('第六章 验证结果与证据范围')
p('本章区分历史运行证据与本次交付回归。历史结果随包保留，本次打包没有重新运行600秒地图试验、120秒对照或完整900秒闭环。')
table(['验证层次', '证据与结果'], [
    ['独立地图场景', '600秒；1171辆插入、817辆到达，0碰撞、0传送'],
    ['后端对照', '两种策略各120秒；任务回传、信号约束与同种子重置一致性通过'],
    ['集成检查', '公共HTTP、页面操作、启动复用、正常停止和重启已检查'],
    ['性能阶段实际闭环', '运行至601.4秒，374辆在场，2847次跨RSU派发，65次信号响应'],
    ['本次交付回归', '前端语法、16项前端、5项后端及7项只读地图一致性检查通过'],
], [1.65, 5.35])
h('120秒对照只说明本次合成负载')
table(['指标', '最早完成', '全部本地'], [
    ['感知任务', '382', '382'],
    ['已完成任务', '205', '189'],
    ['跨站派发', '296', '0'],
    ['未回传任务', '177', '193'],
    ['已完成任务平均时延', '11.97秒', '13.41秒'],
], [3.0, 2.0, 2.0])
p('平均时延只统计已回传任务，不能代表包括未完成项的总体分布。两策略在这次检查中的交通指标相同，没有证明交通效率提升。需要多种子、多车流条件及一致的任务总体统计才能做研究性比较。')
p('601.4秒检查时感知3770项、完成1254项，仍有2516项未回传；完成任务平均时延154.32秒，其中平均排队151.59秒，存在明显积压。当前演示不属于已验证的稳定低时延系统。', bold=True)
p('证据：data/build_stats.json；verification/backend_verification.json；evidence/acceptance.json；evidence/performance/live_smoke.json；evidence/delivery/validation.json。', size=9)

page('第七章 性能设置与问题处理')
p('旧版在暂停或连接中断后仍不断渲染，并逐帧对文字标签写样式、读取尺寸。性能版限制帧率、批量绘制静态物体、减少标签更新、复用并限制特效；静止时停止连续绘制，隐藏页面停止前端状态轮询。')
table(['画质', '帧率上限', '像素比上限', '特效上限'], [
    ['轻量 默认', '24 FPS', '0.85', '12项'],
    ['标准', '30 FPS', '1.00', '24项'],
], [1.9, 1.7, 1.7, 1.7])
h('对照测量的正确读法')
table(['暂停场景指标', '修复前', '性能版'], [
    ['页面主线程忙碌比例', '77.00%', '0.42%'],
    ['观测期间新增3D帧', '3436', '0'],
    ['布局计算次数', '6998', '26'],
    ['最后一帧绘制调用', '125', '43'],
], [3.4, 1.8, 1.8])
p('条件是本机内置Chromium、Intel UHD 730、227辆车固定快照和同一静止视角；旧版观测36.68秒，性能版39.32秒。忙碌比例为TaskDuration除以墙钟时长，不是整机CPU，也不是用户Edge全屏的动态性能结果。')
h('Edge和显卡')
p('本次历史性能测试使用Intel UHD 730，并未直接测量所有Edge环境。硬件加速已开启也可能使用核显；可在edge://gpu查看GL_RENDERER。Windows的应用GPU偏好属于各电脑本地配置，不随ZIP迁移，也不能从设置名称推断实际渲染器已切换。')
h('先按页面状态判断')
p('“已暂停”点继续；“已完成”先导出再重置；“连接中断”先确认启动器和实际端口，再刷新。重置和切换策略会清零当前指标。打开时画面不动，先检查仿真时间与状态，不要连续点击重置。')
p('已知恢复限制：从浏览器后退缓存返回，或图形上下文丢失后若不再同步，用Ctrl+F5重新加载。当前不保证自动恢复图形上下文或浏览器历史缓存中的旧页面。')

page('第八章 可实现能力与后续难度')
p('以下难度是基于当前实现缺口的工程判断，不是工期承诺。三维画面能够展示某一环节，不代表相应的真实感知或硬件能力已经具备。')
table(['层级', '内容', '当前判断与必要条件'], [
    ['已实现', '真实道路与基础三维', 'OSM道路、建筑轮廓、海岸公园、车辆和设备示意可展示'],
    ['已实现', '任务闭环和简单调度', '合成感知、非抢占排队、跨站派发、回传及有界绿灯响应'],
    ['已实现', '交互和记录', '视角、暂停、倍速、策略切换、同种子重置和JSON导出'],
    ['存在难度', '现实交通标定', '需现场车流、车道、转向和配时数据，并校准SUMO行为'],
    ['存在难度', '真实边缘计算接入', '需模型、设备实测、网络接口、资源竞争及结果时效建模'],
    ['存在难度', '稳态低延迟和写实场景', '需负载控制与队列改进；精细建筑与高程需额外数据'],
    ['非常难', '室外多车声学定位', '需同步阵列、多场景标注、噪声与混响处理、多声源分离和泛化验证'],
    ['非常难', '真实信号机闭环与收益', '需安全联锁、故障回退、实地许可及长期对照，不能从动画推导'],
], [1.0, 1.6, 4.4])
h('推荐接入顺序')
p('先核对现场道路与交通需求，再建立录音、阵列位置和车道真值对应的数据集；随后接入已有声学模型，以设备实测替换计算与通信时间；最后在仿真中评估调度和控制策略，再考虑受控的真实系统联调。')
p('评价应同时报告识别或定位误差、端到端时延、未完成与超时比例、队列稳定性和交通指标。不能只比较已完成任务的均值，也不能使用当前理想车道真值作为声学预测结果。')

page('第九章 交付内容和移机说明')
table(['入口或目录', '用途'], [
    ['使用前请先看.txt 和 README.md', '快速启动、地图操作和文件导航'],
    ['docs/', '本手册的可编辑Word版本'],
    ['Start_Demo.* Setup_Runtime.ps1', '启动及固定版本依赖准备；Stop_Demo.*正常停止'],
    ['server.py 和 simulation.py', '接口、仿真、任务调度与控制源代码'],
    ['web/ 和 licenses/', '三维界面、本地依赖与许可'],
    ['data/ 和 scenario/', '原始地图、来源清单、三维场景、网络与合成路线'],
    ['tests/ verification/ evidence/', '检查脚本、历史结果、本次验证与运行档案'],
    ['scripts/ 和 delivery_manifest.json', '重建与打包工具，以及每个文件的SHA-256'],
], [2.65, 4.35])
h('迁移到另一台Windows电脑')
p('Windows x64使用Release便携包可直接解压运行，不需要管理员或修改全局配置。源码包缺少依赖时需要首次联网：启动器从官方源准备Python 3.12.10和SUMO 1.25.0，按固定SHA-256核对并缓存在_runtime。重新运行会复用缓存。')
p('已安装Python 3.10+和SUMO时启动器也可直接复用；下面仅供自定义安装位置。普通运行不需要Node.js或pip安装；生成Word才需要python-docx。源码目录可使用-PrepareOnly预备依赖、-PortableOnly验证便携模式，参数详情见README。')
for line in ("$env:SUMO_HOME = 'D:\\Applications\\Sumo'", "$env:TANGDAO_PYTHON = 'D:\\Applications\\Python311\\python.exe'", "& '.\\Start_Demo.ps1'"):
    x = p(line, size=9); x.paragraph_format.space_after = Pt(2)
h('归档和完整性')
p('示例运行完整JSON保存在evidence/delivery/current_run_export.json；它是运行档案，不是恢复进度的检查点。项目首次启动会新建runs，运行数据不会自动上传。公开包保留源码、地图快照和验证证据，不包含开发机进程文件和个人诊断路径。')
page('第十章 复现与发布核对')
h('文件完整性')
p('完整解压后，在项目根目录执行 python scripts/verify_package.py，逐项核对delivery_manifest.json记录的SHA-256。只有便携Python时，可使用 .\\_runtime\\python-3.12.10\\python.exe scripts/verify_package.py。运行产生的runs和缓存不属于固定交付文件。')
p('runtime-manifest.json记录Python与SUMO的固定版本、官方下载位置和SHA-256。源码首次下载先核对压缩包，再准备运行目录；离线包直接携带准备好的运行时。两个入口均不需要把开发者的Python路径写入系统。')
h('开发回归命令')
table(['命令', '范围'], [
    ['python -m unittest discover -s tests -v', '队列、派发和信号条件单元检查'],
    ['node --check web/app.js', '前端JavaScript语法检查'],
    ['node tests/frontend_performance.mjs', '实际前端函数和Three.js几何；DOM与计时器为测试替身'],
    ['python scripts/verify_package.py', '只读文件哈希核对'],
], [4.0, 3.0])
p('这些检查不能替代真实Edge全屏测量、声学识别实验或多种子交通评估。重跑后端对照会产生新的运行输出，应保存配置、结束时未完成任务数和全部事件，避免只摘取有利均值。')
h('地图重建与发布范围')
p('scripts/build_map.py --verify 会重新生成网络、场景和路线，再执行600秒SUMO；它不是只读命令。普通使用不需要运行。主动增加--download才更新原始OSM，更新后应重新核对信号ID、RSU关联与交通路线。')
p('公开仓库提供源码、数据来源、Word手册与验证证据；大体积Windows便携包在Releases中提供。个人环境路径和原始显卡诊断不公开，项目不会自动上传用户后续产生的运行数据。第三方许可证保存在licenses和_runtime内。')
p('仓库：https://github.com/suxinstar/tangdao-bay-sumo-edge-demo', size=10)
p('发行包：https://github.com/suxinstar/tangdao-bay-sumo-edge-demo/releases', size=10)

out = ROOT / 'docs/唐岛湾声学边缘交通演示_项目交付手册_v1.2.docx'
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
